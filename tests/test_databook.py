from __future__ import annotations

import math
from typing import TYPE_CHECKING

import openpyxl
import pytest

from tag_databook import TagDataBook, defaults
from tag_databook import databook as db

from .conftest import (
    CARBON_FIRST_YEAR,
    CARBON_LAST_YEAR,
    DEFLATOR_FIRST_YEAR,
    DEFLATOR_LAST_YEAR,
    PRICE_YEAR,
    STANDARD_RATES,
    VOT_VALUES,
    build_workbook,
    carbon_row,
    deflator_for,
)

if TYPE_CHECKING:
    import pathlib


class TestValuesOfTime:
    def test_fields_read_from_cells(self, book: TagDataBook) -> None:
        for field, expected in VOT_VALUES.items():
            assert getattr(book.vot, field) == pytest.approx(expected), field

    def test_price_year_is_read_not_assumed(self, book: TagDataBook) -> None:
        assert book.price_year == PRICE_YEAR
        assert book.price_year != defaults.PRICE_YEAR

    def test_accessors(self, book: TagDataBook) -> None:
        assert book.vot_work_rail == book.vot.work_rail
        assert book.vot_commuting == book.vot.commuting
        assert book.vot_other == book.vot.other

    def test_no_fallbacks(self, book: TagDataBook) -> None:
        assert book.fallbacks == ()


class TestDiscountRates:
    def test_bands_use_workbook_rates(self, book: TagDataBook) -> None:
        assert len(book.discount.standard) == len(defaults.DISCOUNT_BANDS)
        for (low, high, rate), (expected_low, expected_high), expected_rate in zip(
            book.discount.standard, defaults.DISCOUNT_BANDS, STANDARD_RATES, strict=True
        ):
            assert (low, high) == (expected_low, expected_high)
            assert rate == pytest.approx(expected_rate)

    def test_rate_lookup(self, book: TagDataBook) -> None:
        assert book.discount_rate_for_year(1) == pytest.approx(STANDARD_RATES[0])
        assert book.discount_rate_for_year(30) == pytest.approx(STANDARD_RATES[0])
        assert book.discount_rate_for_year(31) == pytest.approx(STANDARD_RATES[1])

    def test_rate_past_last_band(self, book: TagDataBook) -> None:
        assert book.discount_rate_for_year(5_000) == pytest.approx(STANDARD_RATES[-1])

    def test_discount_factor_compounds(self, book: TagDataBook) -> None:
        # Inside the first band the stepped factor matches a flat rate.
        assert book.discount_factor(3) == pytest.approx((1 + STANDARD_RATES[0]) ** -3)

    def test_discount_factor_across_bands(self, book: TagDataBook) -> None:
        # The first band covers years 0-30, the second 31-75, so a
        # 40-year factor is 30 years at the first rate and 10 at the
        # second. The fixture's rates rise so each cell can be told
        # apart, the opposite of the real Green Book, so the decline is
        # tested against the real rates in tests/test_defaults.py.
        expected = (1 + STANDARD_RATES[0]) ** -30 * (1 + STANDARD_RATES[1]) ** -10
        assert book.discount_factor(40) == pytest.approx(expected)
        assert book.discount_factor(40) != pytest.approx((1 + STANDARD_RATES[0]) ** -40)

    def test_year_zero_is_undiscounted(self, book: TagDataBook) -> None:
        assert book.discount_factor(0) == pytest.approx(1.0)


class TestCarbonValues:
    def test_published_years_are_read_verbatim(self, book: TagDataBook) -> None:
        for year in (CARBON_FIRST_YEAR, 2025, CARBON_LAST_YEAR):
            low, central, high = carbon_row(year)
            assert book.carbon_price(year, "low") == pytest.approx(low)
            assert book.carbon_price(year, "central") == pytest.approx(central)
            assert book.carbon_price(year, "high") == pytest.approx(high)

    def test_central_is_the_default_scenario(self, book: TagDataBook) -> None:
        assert book.carbon_price(2025) == pytest.approx(carbon_row(2025)[1])

    def test_years_after_series(self, book: TagDataBook) -> None:
        last = carbon_row(CARBON_LAST_YEAR)[1]
        assert book.carbon_price(CARBON_LAST_YEAR + 5) == pytest.approx(last * 1.015**5)

    def test_years_before_series(self, book: TagDataBook) -> None:
        earliest = carbon_row(CARBON_FIRST_YEAR)[1]
        assert book.carbon_price(CARBON_FIRST_YEAR - 10) == pytest.approx(earliest)

    def test_only_the_published_span_is_loaded(self, book: TagDataBook) -> None:
        assert min(book.carbon.factor_cost) == CARBON_FIRST_YEAR
        assert max(book.carbon.factor_cost) == CARBON_LAST_YEAR


class TestDeflator:
    def test_series_spans_the_rows_present(self, book: TagDataBook) -> None:
        assert min(book.gdp.series) == DEFLATOR_FIRST_YEAR
        assert max(book.gdp.series) == DEFLATOR_LAST_YEAR

    def test_deflate_index_ratio(self, book: TagDataBook) -> None:
        expected = 100.0 * deflator_for(2025) / deflator_for(2020)
        assert book.deflate(100.0, 2020, 2025) == pytest.approx(expected)

    def test_deflate_same_year(self, book: TagDataBook) -> None:
        assert book.deflate(250.0, 2021, 2021) == pytest.approx(250.0)

    def test_round_trip_returns_the_original(self, book: TagDataBook) -> None:
        there = book.deflate(80.0, 2016, 2024)
        assert book.deflate(there, 2024, 2016) == pytest.approx(80.0)

    def test_unknown_year_names_the_covered_span(self, book: TagDataBook) -> None:
        with pytest.raises(KeyError, match=r"1990"):
            book.deflate(100.0, 1990, 2020)


class TestValueOfTimeGrowth:
    def test_growth_compounds_at_the_tag_rate(self, book: TagDataBook) -> None:
        grown = book.vot_for_year(10.0, 2020, 2030)
        assert grown == pytest.approx(10.0 * 1.015**10)

    def test_same_year_is_unchanged(self, book: TagDataBook) -> None:
        assert book.vot_for_year(10.0, 2030, 2030) == pytest.approx(10.0)

    def test_going_backwards_is_not_deflated(self, book: TagDataBook) -> None:
        assert book.vot_for_year(10.0, 2030, 2020) == pytest.approx(10.0)


class TestCarValueOfTimeCurve:
    def test_curve_rises_with_distance(self, book: TagDataBook) -> None:
        values = [book.car_vot_by_distance(d) for d in (1, 25, 50, 100, 400)]
        assert values == sorted(values)

    def test_midpoint_is_half_the_asymptote(self, book: TagDataBook) -> None:
        at_midpoint = book.car_vot_by_distance(book.vot.car_xmid)
        assert at_midpoint == pytest.approx(book.vot.car_u_factor_cost / 2.0)

    def test_long_distance(self, book: TagDataBook) -> None:
        assert book.car_vot_by_distance(10_000) == pytest.approx(
            book.vot.car_u_factor_cost, rel=1e-6
        )

    def test_negative_distance(self, book: TagDataBook) -> None:
        # (xmid - d) / k overflows exp() well before this; the guard
        # returns 0.
        assert book.car_vot_by_distance(-1e9) == 0.0
        assert math.isfinite(book.car_vot_by_distance(-1e9))


class TestDegradedWorkbooks:
    def test_moved_cell_recorded(self, tmp_path: pathlib.Path) -> None:
        path = build_workbook(tmp_path / "moved.xlsx")
        workbook = openpyxl.load_workbook(path)
        workbook[db.SHEET_VALUE_OF_TIME][db.VALUE_OF_TIME_CELLS["work_rail"]] = None
        workbook.save(path)

        book = TagDataBook.from_workbook(path)

        assert book.vot.work_rail == pytest.approx(defaults.VALUE_OF_TIME.work_rail)
        assert [f.parameter for f in book.fallbacks] == ["vot.work_rail"]
        assert "work_rail" in str(book.fallbacks[0])

    def test_text_cell_unreadable(self, tmp_path: pathlib.Path) -> None:
        path = build_workbook(tmp_path / "text.xlsx")
        workbook = openpyxl.load_workbook(path)
        workbook[db.SHEET_VALUE_OF_TIME][db.VALUE_OF_TIME_CELLS["commuting"]] = "n/a"
        workbook.save(path)

        book = TagDataBook.from_workbook(path)
        assert [f.parameter for f in book.fallbacks] == ["vot.commuting"]

    def test_missing_sheet_falls_back(self, tmp_path: pathlib.Path) -> None:
        path = build_workbook(tmp_path / "no-vot.xlsx", omit={db.SHEET_VALUE_OF_TIME})
        book = TagDataBook.from_workbook(path)

        assert len(book.fallbacks) == len(db.VALUE_OF_TIME_CELLS)
        assert book.vot == defaults.VALUE_OF_TIME

    def test_no_carbon_sheet(self, tmp_path: pathlib.Path) -> None:
        path = build_workbook(tmp_path / "no-carbon.xlsx", omit={db.SHEET_CARBON})
        book = TagDataBook.from_workbook(path)
        assert book.carbon.factor_cost  # not empty
        assert book.carbon_price(2023) == pytest.approx(defaults.CARBON_CENTRAL_AT_BASE)

    def test_no_deflator_sheet(self, tmp_path: pathlib.Path) -> None:
        path = build_workbook(tmp_path / "no-annual.xlsx", omit={db.SHEET_ANNUAL_PARAMETERS})
        book = TagDataBook.from_workbook(path)

        assert book.gdp.series == {}
        with pytest.raises(KeyError, match="empty"):
            book.deflate(100.0, 2019, 2023)

    def test_unrelated_workbook(self, tmp_path: pathlib.Path) -> None:
        path = tmp_path / "shopping-list.xlsx"
        workbook = openpyxl.Workbook()
        workbook.active["A1"] = "milk"
        workbook.save(path)

        book = TagDataBook.from_workbook(path)
        # Everything falls back, which is exactly what `fallbacks` is
        # for.
        assert len(book.fallbacks) >= len(db.VALUE_OF_TIME_CELLS)
        assert book.vot == defaults.VALUE_OF_TIME


class TestRepr:
    def test_repr_names_the_source(self, book: TagDataBook, workbook_path: pathlib.Path) -> None:
        assert str(workbook_path) in repr(book)
        assert "fallbacks" not in repr(book)

    def test_repr_flags_fallbacks(self, tmp_path: pathlib.Path) -> None:
        path = build_workbook(tmp_path / "partial.xlsx", omit={db.SHEET_VALUE_OF_TIME})
        assert "fallbacks" in repr(TagDataBook.from_workbook(path))
