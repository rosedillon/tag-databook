"""Synthetic TAG Data Book workbooks for the tests.

The real workbook is DfT's to publish, so it is not in the repo. These
fixtures build a workbook with known values at exactly the cells
:mod:`tag_databook.databook` reads, which makes the tests a check on the
layout table itself: move a cell in the parser without moving it here
and the assertions fail.

The values differ from the built-in defaults, so a test can tell a
value read from the workbook from one that fell back to a default.
"""

from __future__ import annotations

from typing import TYPE_CHECKING

import openpyxl
import pytest

from tag_databook import databook as db

if TYPE_CHECKING:
    import pathlib

# Distinct from anything in tag_databook.defaults.
VOT_VALUES: dict[str, float] = {
    "work_car_driver": 11.01,
    "work_car_passenger": 12.02,
    "work_bus": 13.03,
    "work_rail": 14.04,
    "work_walk": 15.05,
    "work_cycle": 16.06,
    "work_average": 17.07,
    "car_u_factor_cost": 40.5,
    "car_xmid": 60.0,
    "car_k": 70.0,
    "commuting": 18.08,
    "other": 19.09,
}

PRICE_YEAR = 2019

STANDARD_RATES = (0.041, 0.042, 0.043, 0.044, 0.045, 0.046)
HEALTH_RATES = (0.051, 0.052, 0.053, 0.054, 0.055, 0.056)
VARIANT_RATES = (0.061, 0.062, 0.063, 0.064, 0.065, 0.066)

CARBON_FIRST_YEAR = 2020
CARBON_LAST_YEAR = 2030
DEFLATOR_FIRST_YEAR = 2015
DEFLATOR_LAST_YEAR = 2025


def carbon_row(year: int) -> tuple[float, float, float]:
    """A recognisable low/central/high triple for *year*."""
    central = 100.0 + (year - CARBON_FIRST_YEAR) * 10.0
    return (central / 2.0, central, central * 2.0)


def deflator_for(year: int) -> float:
    """A deflator index rising 2 points a year, 100 at `PRICE_YEAR`."""
    return 100.0 + (year - PRICE_YEAR) * 2.0


def build_workbook(path: pathlib.Path, *, omit: set[str] | None = None) -> pathlib.Path:
    """Write a synthetic TAG Data Book to *path*.

    Parameters
    ----------
    omit
        Sheet names to leave out entirely, for testing degraded
        workbooks.
    """
    omit = omit or set()
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)

    if db.SHEET_VALUE_OF_TIME not in omit:
        sheet = workbook.create_sheet(db.SHEET_VALUE_OF_TIME)
        sheet[db.PRICE_YEAR_CELL] = PRICE_YEAR
        for field, coordinate in db.VALUE_OF_TIME_CELLS.items():
            sheet[coordinate] = VOT_VALUES[field]

    if db.SHEET_DISCOUNT not in omit:
        sheet = workbook.create_sheet(db.SHEET_DISCOUNT)
        columns = db.DISCOUNT_COLUMNS
        for offset, rates in enumerate(
            zip(STANDARD_RATES, HEALTH_RATES, VARIANT_RATES, strict=True)
        ):
            row = db.DISCOUNT_FIRST_ROW + offset
            standard, health, variant = rates
            sheet[f"{columns['standard']}{row}"] = standard
            sheet[f"{columns['health']}{row}"] = health
            sheet[f"{columns['variant']}{row}"] = variant

    if db.SHEET_CARBON not in omit:
        sheet = workbook.create_sheet(db.SHEET_CARBON)
        sheet[db.PRICE_YEAR_CELL] = PRICE_YEAR
        for offset, year in enumerate(range(CARBON_FIRST_YEAR, CARBON_LAST_YEAR + 1)):
            row = db.CARBON_ROWS.start + offset
            sheet[f"{db.CARBON_YEAR_COLUMN}{row}"] = year
            low, central, high = carbon_row(year)
            for column, value in zip(db.CARBON_VALUE_COLUMNS, (low, central, high), strict=True):
                sheet[f"{column}{row}"] = value

    if db.SHEET_ANNUAL_PARAMETERS not in omit:
        sheet = workbook.create_sheet(db.SHEET_ANNUAL_PARAMETERS)
        for offset, year in enumerate(range(DEFLATOR_FIRST_YEAR, DEFLATOR_LAST_YEAR + 1)):
            row = db.DEFLATOR_ROWS.start + offset
            sheet[f"{db.DEFLATOR_YEAR_COLUMN}{row}"] = year
            sheet[f"{db.DEFLATOR_INDEX_COLUMN}{row}"] = deflator_for(year)

    workbook.save(path)
    return path


@pytest.fixture
def workbook_path(tmp_path: pathlib.Path) -> pathlib.Path:
    """A complete synthetic workbook."""
    return build_workbook(tmp_path / "tag-data-book.xlsx")


@pytest.fixture
def book(workbook_path: pathlib.Path) -> db.TagDataBook:
    """A `TagDataBook` parsed from the complete synthetic workbook."""
    return db.TagDataBook.from_workbook(workbook_path)
